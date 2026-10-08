import io

import pytest
from docx import Document
from openpyxl import Workbook

from app.services.document_extract import extract_document_text


def _pdf_with_text(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Contents 4 0 R "+
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length "+str(len(stream)).encode()+b" >>\nstream\n"+stream+b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = b"%PDF-1.4\n"
    offsets = []
    for number, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(output)
    output += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
    output += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    return output + f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()


def test_extract_docx_paragraphs_and_table_rows():
    document = Document()
    document.add_paragraph("DOCX paragraph")
    cells = document.add_table(rows=1, cols=2).rows[0].cells
    cells[0].text = "left"
    cells[1].text = "right"
    buffer = io.BytesIO()
    document.save(buffer)
    text = extract_document_text(buffer.getvalue(), "docx")
    assert "DOCX paragraph" in text
    assert "left | right" in text


def test_extract_xlsx_rows_and_sheet_boundaries():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "First"
    sheet.append(["Name", "Value"])
    sheet.append(["Boiler", 42])
    workbook.create_sheet("Second").append(["Other", "Row"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    text = extract_document_text(buffer.getvalue(), "xlsx")
    assert "Name | Value" in text
    assert "Boiler | 42" in text
    assert "First" in text and "Second" in text
    assert "Other | Row" in text


@pytest.mark.parametrize("kind", ["txt", "md"])
def test_extract_plain_text_decodes_utf8_with_replacement(kind):
    assert extract_document_text(b"hello\xff", kind) == "hello\ufffd"


def test_extract_csv_uses_row_grouped_text():
    text = extract_document_text(b"name,value\nBoiler,42\n", "csv")
    assert "name: Boiler" in text
    assert "value: 42" in text


def test_extract_pdf_embedded_text():
    text = extract_document_text(_pdf_with_text("embedded PDF text"), "pdf")
    assert "embedded PDF text" in text


def test_extract_image_returns_empty_text():
    assert extract_document_text(b"image bytes", "image") == ""
