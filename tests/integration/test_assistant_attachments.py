from pathlib import Path

import pytest
import httpx
import respx
from app.db import SessionLocal
from app.config import get_settings
from app.models import AssistantAttachment, AuditLog, InternalChatThread, Project, ProjectFile, User
from app.services.users import hash_password

from app.chat.assistant.attachments import (
    MAX_ATTACHMENTS_PER_THREAD,
    MAX_FILE_SIZE,
    allowed_attachment,
    build_multimodal_content,
    cleanup_attachments,
    prepare_attachment_context,
)


def test_allowed_attachment_accepts_supported_document_and_rejects_mismatch():
    assert allowed_attachment("manual.PDF", "application/pdf", 120) == "pdf"
    assert allowed_attachment("data.csv", "text/csv", 120) == "csv"
    with pytest.raises(ValueError, match="unsupported"):
        allowed_attachment("manual.pdf", "text/plain", 120)


def test_allowed_attachment_accepts_docx_xlsx_and_markdown():
    docx = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    assert allowed_attachment("doc.docx", docx, 120) == "docx"
    assert allowed_attachment("sheet.xlsx", xlsx, 120) == "xlsx"
    assert allowed_attachment("notes.md", "text/markdown", 120) == "md"
    assert allowed_attachment("notes.md", "text/plain", 120) == "md"


def test_document_context_extracts_real_docx(tmp_path):
    from docx import Document as DocxDocument

    document = DocxDocument()
    document.add_paragraph("Hola desde docx")
    source = tmp_path / "a.docx"
    document.save(str(source))
    attachment = type("Attachment", (), {"kind": "docx", "stored_path": str(source), "filename": "a.docx", "id": 11})()

    result = prepare_attachment_context([attachment], max_chars=1000)

    assert result[0]["kind"] == "docx"
    assert "Hola desde docx" in result[0]["text"]


def test_document_context_extracts_real_xlsx(tmp_path):
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.append(["producto", "plan"])
    sheet.append(["internet", "fibra 1000"])
    source = tmp_path / "a.xlsx"
    book.save(str(source))
    attachment = type("Attachment", (), {"kind": "xlsx", "stored_path": str(source), "filename": "a.xlsx", "id": 12})()

    result = prepare_attachment_context([attachment], max_chars=1000)

    assert result[0]["kind"] == "xlsx"
    assert "fibra 1000" in result[0]["text"]


def test_attachment_limits_cover_size_and_count():
    with pytest.raises(ValueError, match="size"):
        allowed_attachment("photo.png", "image/png", MAX_FILE_SIZE + 1)
    assert MAX_ATTACHMENTS_PER_THREAD == 10


def test_document_context_uses_extracted_text_and_is_bounded(monkeypatch, tmp_path):
    monkeypatch.setattr(
        "app.chat.assistant.attachments.extract_document_text",
        lambda content, kind: "manual: " + content.decode(),
    )
    source = tmp_path / "a.txt"
    source.write_bytes(b"hello")
    attachment = type("Attachment", (), {"kind": "txt", "stored_path": str(source), "filename": "a.txt"})()
    result = prepare_attachment_context([attachment], max_chars=12)
    assert result == [{"type": "document", "filename": "a.txt", "kind": "txt", "text": "manual: hell", "attachment_id": None}]


def test_limits_are_read_from_settings(monkeypatch):
    monkeypatch.setenv("ASSISTANT_ATTACHMENT_MAX_BYTES", "3")
    get_settings.cache_clear()
    with pytest.raises(ValueError, match="size"):
        allowed_attachment("a.txt", "text/plain", 4)
    monkeypatch.setenv("ASSISTANT_ATTACHMENT_MAX_BYTES", "5242880")
    get_settings.cache_clear()


def test_document_context_preserves_source_identity_in_prompt(tmp_path):
    source = tmp_path / "manual.txt"
    source.write_bytes(b"source text")
    attachment = type("Attachment", (), {"kind": "txt", "stored_path": str(source), "filename": "manual.txt", "id": 42})()
    result = prepare_attachment_context([attachment], max_chars=100)
    assert result[0]["filename"] == "manual.txt"
    assert result[0]["attachment_id"] == 42
    assert result[0]["kind"] == "txt"


def test_images_are_openai_parts_and_oversized_images_are_rejected(tmp_path):
    image = tmp_path / "photo.png"
    image.write_bytes(b"png")
    attachment = type("Attachment", (), {"kind": "image", "stored_path": str(image), "filename": "photo.png", "id": 7})()
    parts = build_multimodal_content([attachment], "Describe this")
    assert parts[0] == {"type": "text", "text": "Describe this"}
    assert parts[1]["type"] == "image_url"
    assert parts[1]["image_url"]["url"].startswith("data:image/png;base64,")


def test_cleanup_removes_only_orphaned_assistant_files(tmp_path):
    orphan = tmp_path / "orphan"
    active = tmp_path / "active"
    orphan.write_bytes(b"x")
    active.write_bytes(b"x")
    rows = [
        type("Attachment", (), {"stored_path": str(orphan), "thread_id": 1, "deleted_at": None})(),
        type("Attachment", (), {"stored_path": str(active), "thread_id": 2, "deleted_at": None})(),
    ]
    assert cleanup_attachments(rows, active_thread_ids={2}) == 1
    assert not orphan.exists()
    assert active.exists()


def test_cleanup_removes_expired_attachment_even_on_active_thread(tmp_path):
    expired = tmp_path / "expired"
    expired.write_bytes(b"x")
    row = type("Attachment", (), {"stored_path": str(expired), "thread_id": 2, "deleted_at": None})()
    assert cleanup_attachments([row], active_thread_ids={2}, expired=True) == 1
    assert not expired.exists()


def test_admin_upload_persists_private_thread_scoped_reference(client, monkeypatch, tmp_path):
    username = "attachment-admin"
    with SessionLocal() as session:
        session.add(User(username=username, password_hash=hash_password("pass"), role="admin"))
        session.commit()
    assert client.post("/panel/api/login", json={"username": username, "password": "pass"}).status_code == 200
    thread_id = client.post("/panel/api/chat/threads", json={"proyecto": "default"}).json()["thread"]["id"]
    monkeypatch.setenv("MEDIA_DIR", str(tmp_path))
    get_settings.cache_clear()
    response = client.post(f"/panel/api/chat/threads/{thread_id}/attachments", files={"file": ("notes.txt", b"private notes", "text/plain")})
    assert response.status_code == 201, response.text
    payload = response.json()["attachment"]
    assert payload["filename"] == "notes.txt"
    assert "stored_path" not in payload
    assert list(tmp_path.glob(f"assistant/{thread_id}/*"))


def test_upload_rejection_paths_return_errors_without_success_audit(client, monkeypatch):
    username = "attachment-rejections"
    with SessionLocal() as session:
        session.add(User(username=username, password_hash=hash_password("pass"), role="admin"))
        session.commit()
    assert client.post("/panel/api/login", json={"username": username, "password": "pass"}).status_code == 200
    thread_id = client.post("/panel/api/chat/threads", json={"proyecto": "default"}).json()["thread"]["id"]
    with SessionLocal() as session:
        before = session.query(AuditLog).filter(AuditLog.action == "internal_chat.attachment_uploaded").count()
    bad = client.post(f"/panel/api/chat/threads/{thread_id}/attachments", files={"file": ("x.exe", b"x", "application/octet-stream")})
    assert bad.status_code == 400
    monkeypatch.setenv("ASSISTANT_ATTACHMENT_MAX_BYTES", "2")
    get_settings.cache_clear()
    oversized = client.post(f"/panel/api/chat/threads/{thread_id}/attachments", files={"file": ("x.txt", b"large", "text/plain")})
    assert oversized.status_code == 400
    with SessionLocal() as session:
        assert session.query(AuditLog).filter(AuditLog.action == "internal_chat.attachment_uploaded").count() == before
    get_settings.cache_clear()


def test_cross_thread_attachment_upload_is_denied(client):
    first, second = "attachment-owner-a", "attachment-owner-b"
    with SessionLocal() as session:
        session.add_all([User(username=first, password_hash=hash_password("pass"), role="admin"), User(username=second, password_hash=hash_password("pass"), role="admin")])
        session.commit()
    assert client.post("/panel/api/login", json={"username": first, "password": "pass"}).status_code == 200
    own = client.post("/panel/api/chat/threads", json={"proyecto": "default"}).json()["thread"]["id"]
    client.post("/panel/api/logout")
    assert client.post("/panel/api/login", json={"username": second, "password": "pass"}).status_code == 200
    denied = client.post(f"/panel/api/chat/threads/{own}/attachments", files={"file": ("x.txt", b"x", "text/plain")})
    assert denied.status_code == 404


def test_attachment_delete_unlinks_file_and_hides_metadata(client, monkeypatch, tmp_path):
    username = "attachment-delete-owner"
    with SessionLocal() as session:
        session.add(User(username=username, password_hash=hash_password("pass"), role="admin"))
        session.commit()
    assert client.post("/panel/api/login", json={"username": username, "password": "pass"}).status_code == 200
    thread_id = client.post("/panel/api/chat/threads", json={"proyecto": "default"}).json()["thread"]["id"]
    monkeypatch.setenv("MEDIA_DIR", str(tmp_path))
    get_settings.cache_clear()
    uploaded = client.post(f"/panel/api/chat/threads/{thread_id}/attachments", files={"file": ("notes.txt", b"private notes", "text/plain")})
    attachment_id = uploaded.json()["attachment"]["id"]
    stored_file = next((tmp_path / "assistant" / str(thread_id)).iterdir())

    deleted = client.delete(f"/panel/api/chat/threads/{thread_id}/attachments/{attachment_id}")

    assert deleted.status_code == 204
    assert deleted.content == b""
    assert not stored_file.exists()
    assert client.get(f"/panel/api/chat/threads/{thread_id}").json()["attachments"] == []
    assert client.delete(f"/panel/api/chat/threads/{thread_id}/attachments/{attachment_id}").status_code == 404
    get_settings.cache_clear()


def test_attachment_delete_is_thread_and_owner_scoped_and_guarded(client, monkeypatch, tmp_path):
    owners = ["attachment-delete-a", "attachment-delete-b"]
    with SessionLocal() as session:
        session.add_all(User(username=name, password_hash=hash_password("pass"), role="admin") for name in owners)
        session.commit()
    client.post("/panel/api/login", json={"username": owners[0], "password": "pass"})
    own_thread = client.post("/panel/api/chat/threads", json={"proyecto": "default"}).json()["thread"]["id"]
    other_thread = client.post("/panel/api/chat/threads", json={"proyecto": "default"}).json()["thread"]["id"]
    monkeypatch.setenv("MEDIA_DIR", str(tmp_path))
    get_settings.cache_clear()
    attachment_id = client.post(f"/panel/api/chat/threads/{other_thread}/attachments", files={"file": ("x.txt", b"x", "text/plain")}).json()["attachment"]["id"]
    assert client.delete(f"/panel/api/chat/threads/{own_thread}/attachments/{attachment_id}").status_code == 404
    with SessionLocal() as session:
        thread = session.get(InternalChatThread, other_thread)
        thread.assistant_state = {"live": {"status": "running"}}
        session.commit()
    assert client.delete(f"/panel/api/chat/threads/{other_thread}/attachments/{attachment_id}").status_code == 409
    client.post("/panel/api/logout")
    client.post("/panel/api/login", json={"username": owners[1], "password": "pass"})
    assert client.delete(f"/panel/api/chat/threads/{other_thread}/attachments/{attachment_id}").status_code == 404
    get_settings.cache_clear()


def test_attachment_content_is_inline_owner_scoped_and_missing_file_is_404(client, monkeypatch, tmp_path):
    username = "attachment-content-owner"
    with SessionLocal() as session:
        user = User(username=username, password_hash=hash_password("pass"), role="admin")
        session.add(user)
        session.flush()
        owner_id = user.id
        session.add(User(username="attachment-content-other", password_hash=hash_password("pass"), role="admin"))
        session.commit()
    client.post("/panel/api/login", json={"username": username, "password": "pass"})
    thread_id = client.post("/panel/api/chat/threads", json={"proyecto": "default"}).json()["thread"]["id"]
    monkeypatch.setenv("MEDIA_DIR", str(tmp_path))
    get_settings.cache_clear()
    uploaded = client.post(f"/panel/api/chat/threads/{thread_id}/attachments", files={"file": ("photo.png", b"png-bytes", "image/png")})
    attachment_id = uploaded.json()["attachment"]["id"]
    content_url = f"/panel/api/chat/threads/{thread_id}/attachments/{attachment_id}/content"
    content = client.get(content_url)
    assert content.status_code == 200
    assert content.content == b"png-bytes"
    assert content.headers["content-type"] == "image/png"
    assert content.headers["content-disposition"].startswith('inline; filename="photo.png"')
    client.post("/panel/api/logout")
    client.post("/panel/api/login", json={"username": "attachment-content-other", "password": "pass"})
    assert client.get(content_url).status_code == 404
    client.post("/panel/api/logout")
    client.post("/panel/api/login", json={"username": username, "password": "pass"})
    with SessionLocal() as session:
        row = session.get(AssistantAttachment, attachment_id)
        Path(row.stored_path).unlink()
    assert client.get(content_url).status_code == 404
    client.delete(f"/panel/api/chat/threads/{thread_id}/attachments/{attachment_id}")
    assert client.get(content_url).status_code == 404
    get_settings.cache_clear()


def _project_file(session, project_id, owner_id, path, *, filename="notes.txt", content_type="text/plain"):
    path.write_bytes(b"project notes")
    row = ProjectFile(
        project_id=project_id, owner_id=owner_id, filename=filename, content_type=content_type,
        size=len(b"project notes"), sha256="a" * 64, stored_path=str(path), status="ready",
    )
    session.add(row)
    session.flush()
    return row.id


def test_project_file_can_be_attached_to_its_project_thread(client, tmp_path, monkeypatch):
    username = "project-attachment-owner"
    with SessionLocal() as session:
        user = User(username=username, password_hash=hash_password("pass"), role="admin")
        session.add(user)
        session.flush()
        project = Project(slug="project-attachment", name="Project", purpose="A", purpose_embedding=[0.0] * 384, owner_id=user.id)
        session.add(project)
        session.flush()
        project_id = project.id
        source_path = tmp_path / "source.txt"
        file_id = _project_file(session, project_id, user.id, source_path)
        session.commit()
    assert client.post("/panel/api/login", json={"username": username, "password": "pass"}).status_code == 200
    thread_id = client.post("/panel/api/chat/threads", json={"proyecto": "project-attachment"}).json()["thread"]["id"]
    monkeypatch.setenv("MEDIA_DIR", str(tmp_path / "media"))
    get_settings.cache_clear()

    response = client.post(
        f"/panel/api/chat/threads/{thread_id}/attachments/from-project", json={"project_file_id": file_id}
    )

    assert response.status_code == 201, response.text
    assert response.json()["attachment"]["filename"] == "notes.txt"
    assert list((tmp_path / "media").glob(f"assistant/{thread_id}/*"))


def test_project_file_attachment_rejects_other_project_and_unsupported_type(client, tmp_path):
    username = "project-attachment-boundary"
    with SessionLocal() as session:
        user = User(username=username, password_hash=hash_password("pass"), role="admin")
        session.add(user)
        session.flush()
        first = Project(slug="project-first", name="First", purpose="A", purpose_embedding=[0.0] * 384, owner_id=user.id)
        second = Project(slug="project-second", name="Second", purpose="A", purpose_embedding=[0.0] * 384, owner_id=user.id)
        session.add_all([first, second])
        session.flush()
        foreign_id = _project_file(session, second.id, user.id, tmp_path / "foreign.txt")
        unsupported_id = _project_file(
            session, first.id, user.id, tmp_path / "unsupported.zip", filename="notes.zip", content_type="application/zip"
        )
        session.commit()
    assert client.post("/panel/api/login", json={"username": username, "password": "pass"}).status_code == 200
    thread_id = client.post("/panel/api/chat/threads", json={"proyecto": "project-first"}).json()["thread"]["id"]

    foreign = client.post(
        f"/panel/api/chat/threads/{thread_id}/attachments/from-project", json={"project_file_id": foreign_id}
    )
    unsupported = client.post(
        f"/panel/api/chat/threads/{thread_id}/attachments/from-project", json={"project_file_id": unsupported_id}
    )

    assert foreign.status_code == 404
    assert unsupported.status_code == 400
    assert "unsupported attachment type" in unsupported.json()["detail"]


def test_project_file_attachment_obeys_count_limit_and_default_thread_has_no_project(client, monkeypatch, tmp_path):
    username = "project-attachment-limit"
    with SessionLocal() as session:
        user = User(username=username, password_hash=hash_password("pass"), role="admin")
        session.add(user)
        session.flush()
        project = Project(slug="project-limit", name="Project", purpose="A", purpose_embedding=[0.0] * 384, owner_id=user.id)
        session.add(project)
        session.flush()
        file_id = _project_file(session, project.id, user.id, tmp_path / "limit.txt")
        session.commit()
    assert client.post("/panel/api/login", json={"username": username, "password": "pass"}).status_code == 200
    project_thread = client.post("/panel/api/chat/threads", json={"proyecto": "project-limit"}).json()["thread"]["id"]
    default_thread = client.post("/panel/api/chat/threads", json={"proyecto": "default"}).json()["thread"]["id"]
    monkeypatch.setenv("ASSISTANT_ATTACHMENT_MAX_COUNT", "0")
    get_settings.cache_clear()

    limited = client.post(
        f"/panel/api/chat/threads/{project_thread}/attachments/from-project", json={"project_file_id": file_id}
    )
    no_project = client.post(
        f"/panel/api/chat/threads/{default_thread}/attachments/from-project", json={"project_file_id": file_id}
    )

    assert limited.status_code == 400
    assert "count limit" in limited.json()["detail"]
    assert no_project.status_code == 400
    get_settings.cache_clear()


def test_expired_cleanup_deletes_metadata_for_active_thread(tmp_path):
    from datetime import datetime, timedelta, timezone
    from app.chat.assistant.attachments import cleanup_expired
    path = tmp_path / "expired"
    path.write_bytes(b"x")
    with SessionLocal() as session:
        user = User(username="cleanup-owner", password_hash=hash_password("pass"), role="admin")
        session.add(user)
        session.flush()
        thread = InternalChatThread(user_id=user.id)
        session.add(thread)
        session.flush()
        row = AssistantAttachment(thread_id=thread.id, filename="x.txt", kind="txt", content_type="text/plain", size=1, sha256="x", stored_path=str(path), created_at=datetime.now(timezone.utc) - timedelta(days=40))
        session.add(row)
        session.flush()
        assert cleanup_expired(session, datetime.now(timezone.utc) - timedelta(days=30), {thread.id}) == 1
        session.commit()
        assert session.get(AssistantAttachment, row.id) is None
    assert not path.exists()


def test_multimodal_invocation_uses_openrouter_image_part_and_audits(tmp_path, monkeypatch):
    image = tmp_path / "photo.png"
    image.write_bytes(b"png")
    attachment = type("Attachment", (), {"kind": "image", "stored_path": str(image), "filename": "photo.png", "id": 77})()
    monkeypatch.setenv("OPENROUTER_API_KEY", "sk-test")
    get_settings.cache_clear()
    with respx.mock(base_url="https://openrouter.ai/api/v1") as router:
        route = router.post("/chat/completions").mock(
            return_value=httpx.Response(200, headers={"content-type": "text/event-stream"}, content='data: {"choices":[{"delta":{"content":"Seen"}}]}\n\ndata: [DONE]\n\n'))
        from app.chat.assistant.nodes import execute
        with SessionLocal() as session:
            result = execute({"thread_id": 9, "run_id": "run-image", "content": "Describe", "plan": {"tool": None, "arguments": {}, "requires_approval": False}, "attachments": [attachment], "on_token": lambda token: None, "session": session, "actor": "attachment-admin"})
            session.flush()
            actions = {row.action for row in session.query(AuditLog).filter(AuditLog.entity_id == "9").all()}
            session.rollback()
        assert result["status"] == "completed"
        assert "internal_chat.attachments_used" in actions
        assert route.called
        request = route.calls[0].request
        body = __import__("json").loads(request.content)
        assert body["model"] == "google/gemini-2.5-flash"
        assert body["messages"][1]["content"][1]["type"] == "image_url"
    get_settings.cache_clear()
