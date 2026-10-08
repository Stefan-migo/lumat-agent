"""Text extraction for project-file formats."""

import io

from app.services.kb import _csv_rows_to_text


def extract_document_text(content: bytes, kind: str) -> str:
    """Extract searchable text from an uploaded project file."""
    if kind == "pdf":
        import pdfplumber

        with pdfplumber.open(io.BytesIO(content)) as pdf:
            pages = []
            for page in pdf.pages:
                blocks = [page.extract_text() or ""]
                blocks.extend(" | ".join(cell or "" for cell in row) for table in page.extract_tables() for row in table)
                pages.append("\n".join(block for block in blocks if block))
            return "\n".join(pages)
    if kind == "docx":
        from docx import Document

        document = Document(io.BytesIO(content))
        blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
        blocks.extend(
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        )
        return "\n".join(blocks)
    if kind == "xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        blocks = []
        try:
            for sheet in workbook:
                rows = [
                    " | ".join(str(value) if value is not None else "" for value in row)
                    for row in sheet.iter_rows(values_only=True)
                ]
                rows = [row for row in rows if row.strip(" |")]
                if rows:
                    blocks.extend((f"Sheet: {sheet.title}", *rows))
        finally:
            workbook.close()
        return "\n".join(blocks)
    if kind in {"txt", "md"}:
        return content.decode("utf-8", errors="replace")
    if kind == "csv":
        return _csv_rows_to_text(content.decode("utf-8", errors="replace"))
    if kind == "image":
        return ""
    raise ValueError(f"unsupported project file kind: {kind}")
